"""
Converts Research_Paper_Single_Column.md directly to Research_Paper_Single_Column.docx
applying Times New Roman, 12pt, 1.15 line spacing, 1-inch margins, and styled tables.
"""

import re
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn


def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>\n'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>\n'
        f'  <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)


def set_cell_shading(cell, color_hex="F2F2F2"):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def convert(md_path: Path, docx_path: Path):
    doc = docx.Document()

    # 1. Page Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        # Single column layout
        sectPr = section._sectPr
        cols = sectPr.xpath("./w:cols")
        if cols:
            cols[0].set(qn("w:num"), "1")

    # 2. Normal Style
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)
    normal_style.paragraph_format.space_before = Pt(0)

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    i = 0
    in_code_block = False
    in_table = False
    table_lines = []

    while i < len(lines):
        line = lines[i].rstrip("\r\n")

        # Skip horizontal rules
        if line.strip() == "---":
            i += 1
            continue

        # Code block handling
        if line.startswith("```"):
            in_code_block = not in_code_block
            i += 1
            continue

        if in_code_block:
            # Add mono/diagram paragraph
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line)
            run.font.name = "Courier New"
            run.font.size = Pt(8.5)
            i += 1
            continue

        # Table handling
        if line.startswith("|") and line.endswith("|"):
            table_lines.append(line)
            i += 1
            # Check if table continues
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            # Process table
            if len(table_lines) >= 2:
                # filter out separator line
                parsed_rows = []
                for t_line in table_lines:
                    if re.match(r"^\|[\s\-:]+\|$", t_line):
                        continue
                    cells = [c.strip() for c in t_line.strip("|").split("|")]
                    parsed_rows.append(cells)

                if parsed_rows:
                    num_cols = max(len(r) for r in parsed_rows)
                    tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

                    for row_idx, r_data in enumerate(parsed_rows):
                        is_header = (row_idx == 0)
                        for col_idx, val in enumerate(r_data):
                            if col_idx < num_cols:
                                cell = tbl.cell(row_idx, col_idx)
                                # strip markdown bold
                                clean_val = re.sub(r"\*\*(.*?)\*\*", r"\1", val)
                                cell.text = clean_val
                                if is_header:
                                    set_cell_shading(cell, "E0E0E0")
                                elif row_idx % 2 == 1:
                                    set_cell_shading(cell, "F9F9F9")
                                set_cell_border(cell)
                                p = cell.paragraphs[0]
                                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx in [0, 1, 2] else WD_ALIGN_PARAGRAPH.CENTER
                                p.paragraph_format.line_spacing = 1.0
                                p.paragraph_format.space_after = Pt(2)
                                for r in p.runs:
                                    r.font.name = "Times New Roman"
                                    r.font.size = Pt(10 if is_header else 9.5)
                                    if is_header or "**" in val:
                                        r.bold = True
            table_lines = []
            continue

        # Title (# )
        if line.startswith("# "):
            title_text = line[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(title_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(16)
            run.bold = True
            i += 1
            continue

        # Heading 1 (## )
        if line.startswith("## "):
            h_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(h_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
            run.bold = True
            i += 1
            continue

        # Heading 2 (### )
        if line.startswith("### "):
            h_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(h_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12.5)
            run.bold = True
            i += 1
            continue

        # Heading 3 (#### )
        if line.startswith("#### "):
            h_text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(h_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.bold = True
            i += 1
            continue

        # Empty lines
        if not line.strip():
            i += 1
            continue

        # Regular paragraphs and metadata
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.space_before = Pt(0)

        # Handle bold/italic markdown spans in line
        # Simple parser for bold **text** and *italic*
        tokens = re.split(r"(\*\*.*?\*\*|\*.*?\*)", line)
        for token in tokens:
            if token.startswith("**") and token.endswith("**"):
                r = p.add_run(token[2:-2])
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
                r.bold = True
            elif token.startswith("*") and token.endswith("*"):
                r = p.add_run(token[1:-1])
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
                r.italic = True
            else:
                r = p.add_run(token)
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)

        # Center author metadata
        if line.startswith("**Author:**") or line.startswith("**Affiliation:**") or line.startswith("**Document Style:**"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        i += 1

    try:
        doc.save(str(docx_path))
        print(f"Successfully converted {md_path} -> {docx_path}")
    except PermissionError:
        fallback_path = docx_path.with_name(f"{docx_path.stem}_Updated.docx")
        doc.save(str(fallback_path))
        print(f"Notice: '{docx_path.name}' is currently open in Microsoft Word.")
        print(f"Saved updated file to: {fallback_path}")


if __name__ == "__main__":
    convert(Path("Research_Paper_Single_Column.md"), Path("Research_Paper_Single_Column.docx"))
